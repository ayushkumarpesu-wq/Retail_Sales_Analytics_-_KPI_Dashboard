"""Clean UCI Online Retail data and build CSV, Excel, and SQLite outputs."""

from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE_FILE = ROOT / "data" / "raw" / "Online Retail.xlsx"
OUTPUT_DIR = ROOT / "data" / "processed"
SAMPLE_DIR = ROOT / "data" / "sample"
REPORT_DIR = ROOT / "reports"
OUTPUT_CSV = OUTPUT_DIR / "retail_sales_clean.csv"
SAMPLE_CSV = SAMPLE_DIR / "retail_sales_sample.csv"
OUTPUT_XLSX = OUTPUT_DIR / "retail_sales_clean.xlsx"
OUTPUT_DB = OUTPUT_DIR / "retail_sales.db"
TABLE_NAME = "retail_sales"

COLUMN_MAP = {
    "InvoiceNo": "invoice_no",
    "StockCode": "stock_code",
    "Description": "description",
    "Quantity": "quantity",
    "InvoiceDate": "invoice_date",
    "UnitPrice": "unit_price",
    "CustomerID": "customer_id",
    "Country": "country",
}
REQUIRED_COLUMNS = set(COLUMN_MAP)


def _outlier_bounds(values: pd.Series) -> tuple[float, float]:
    q1 = values.quantile(0.25)
    q3 = values.quantile(0.75)
    iqr = q3 - q1
    return q1 - 1.5 * iqr, q3 + 1.5 * iqr


def clean_data(raw: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, int]]:
    """Return cleaned transactions and counts explaining each cleaning step."""
    missing_columns = REQUIRED_COLUMNS.difference(raw.columns)
    if missing_columns:
        raise ValueError(f"Dataset is missing expected columns: {sorted(missing_columns)}")

    starting_rows = len(raw)
    data = raw.rename(columns=COLUMN_MAP).copy()
    data = data[list(COLUMN_MAP.values())]

    for column in ("invoice_no", "stock_code", "description", "country"):
        data[column] = data[column].astype("string").str.strip()
    data["invoice_no"] = data["invoice_no"].replace("", pd.NA)
    data["stock_code"] = data["stock_code"].replace("", pd.NA)
    data["description"] = data["description"].replace("", pd.NA).fillna("Unknown")
    data["country"] = data["country"].replace("", pd.NA)
    data["invoice_date"] = pd.to_datetime(data["invoice_date"], errors="coerce")
    data["quantity"] = pd.to_numeric(data["quantity"], errors="coerce")
    data["unit_price"] = pd.to_numeric(data["unit_price"], errors="coerce")
    data["customer_id"] = pd.to_numeric(data["customer_id"], errors="coerce").astype("Int64")

    before_deduplication = len(data)
    data = data.drop_duplicates()
    duplicate_rows = before_deduplication - len(data)

    valid = (
        data["invoice_no"].notna()
        & data["stock_code"].notna()
        & data["invoice_date"].notna()
        & data["country"].notna()
        & data["quantity"].notna()
        & data["unit_price"].notna()
        & data["quantity"].ne(0)
        & data["unit_price"].ge(0)
    )
    invalid_rows = int((~valid).sum())
    data = data.loc[valid].copy()

    data["is_cancellation"] = data["invoice_no"].str.upper().str.startswith("C")
    data["is_return"] = (data["quantity"] < 0) & ~data["is_cancellation"]
    data["revenue"] = data["quantity"] * data["unit_price"]
    data["invoice_month"] = data["invoice_date"].dt.to_period("M").astype(str)

    positive_sales = data.loc[data["quantity"] > 0]
    quantity_low, quantity_high = _outlier_bounds(positive_sales["quantity"])
    price_low, price_high = _outlier_bounds(positive_sales["unit_price"])
    data["quantity_outlier"] = (
        (data["quantity"] > 0)
        & ~data["quantity"].between(quantity_low, quantity_high, inclusive="both")
    )
    data["unit_price_outlier"] = (
        (data["quantity"] > 0)
        & ~data["unit_price"].between(price_low, price_high, inclusive="both")
    )
    data["is_outlier"] = data["quantity_outlier"] | data["unit_price_outlier"]

    data = data.sort_values(["invoice_date", "invoice_no", "stock_code"]).reset_index(drop=True)
    quality = {
        "source_rows": starting_rows,
        "duplicate_rows_removed": duplicate_rows,
        "invalid_rows_removed": invalid_rows,
        "clean_rows": len(data),
        "missing_customer_id_rows_retained": int(data["customer_id"].isna().sum()),
        "return_rows_retained": int(data["is_return"].sum()),
        "cancellation_rows_retained": int(data["is_cancellation"].sum()),
        "outlier_rows_flagged_not_removed": int(data["is_outlier"].sum()),
    }
    return data, quality


def write_outputs(data: pd.DataFrame, quality: dict[str, int]) -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    SAMPLE_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    data.to_csv(OUTPUT_CSV, index=False, date_format="%Y-%m-%d %H:%M:%S")
    sample = data.sample(n=min(500, len(data)), random_state=42).sort_values(
        ["invoice_date", "invoice_no", "stock_code"]
    )
    sample.to_csv(SAMPLE_CSV, index=False, date_format="%Y-%m-%d %H:%M:%S")

    monthly = (
        data.groupby("invoice_month", as_index=False)["revenue"]
        .sum()
        .rename(columns={"revenue": "net_revenue"})
    )
    completed = data.loc[~data["is_cancellation"] & (data["quantity"] > 0)]
    orders_by_country = completed.groupby("country")["invoice_no"].nunique()
    countries = (
        data.groupby("country")["revenue"]
        .sum()
        .rename("net_revenue")
        .to_frame()
        .join(orders_by_country.rename("completed_orders"))
        .fillna({"completed_orders": 0})
        .reset_index()
        .sort_values("net_revenue", ascending=False)
    )
    products = (
        data.groupby(["stock_code", "description"], as_index=False)
        .agg(net_revenue=("revenue", "sum"), net_quantity=("quantity", "sum"))
        .sort_values("net_revenue", ascending=False)
        .head(100)
    )

    workbook = Workbook(write_only=True)
    for name, frame in (
        ("Monthly", monthly),
        ("Countries", countries),
        ("Top Products", products),
    ):
        sheet = workbook.create_sheet(name)
        sheet.append(list(frame.columns))
        for row in frame.itertuples(index=False, name=None):
            sheet.append(
                [
                    None if pd.isna(value) else value.item()
                    if hasattr(value, "item") else value
                    for value in row
                ]
            )
    quality_sheet = workbook.create_sheet("Data Quality")
    quality_sheet.append(["metric", "value"])
    for metric, value in quality.items():
        quality_sheet.append([metric, value])
    workbook.save(OUTPUT_XLSX)

    with sqlite3.connect(OUTPUT_DB) as connection:
        data.to_sql(TABLE_NAME, connection, if_exists="replace", index=False)
        connection.execute(
            f"CREATE INDEX IF NOT EXISTS idx_{TABLE_NAME}_month "
            f"ON {TABLE_NAME}(invoice_month)"
        )
        connection.execute(
            f"CREATE INDEX IF NOT EXISTS idx_{TABLE_NAME}_country "
            f"ON {TABLE_NAME}(country)"
        )
        connection.commit()

    (REPORT_DIR / "data_quality.json").write_text(
        json.dumps(quality, indent=2), encoding="utf-8"
    )


def write_findings_report(data: pd.DataFrame, quality: dict[str, int]) -> None:
    """Create a concise report whose claims are calculated from the cleaned data."""
    completed = data.loc[~data["is_cancellation"] & (data["quantity"] > 0)]
    completed_invoice_ids = completed["invoice_no"].unique()
    completed_revenue = data.loc[
        ~data["is_cancellation"] & data["invoice_no"].isin(completed_invoice_ids),
        "revenue",
    ].sum()
    net_revenue = data["revenue"].sum()
    average_order_value = completed_revenue / len(completed_invoice_ids) if len(completed_invoice_ids) else 0

    country_sales = data.groupby("country")["revenue"].sum().sort_values(ascending=False)
    monthly_sales = data.groupby("invoice_month")["revenue"].sum().sort_values(ascending=False)
    product_sales = (
        data.groupby(["stock_code", "description"])["revenue"]
        .sum()
        .sort_values(ascending=False)
    )
    return_rows = int(data["is_return"].sum())
    cancellation_rows = int(data["is_cancellation"].sum())
    return_revenue = data.loc[data["is_return"], "revenue"].sum()
    top_country, top_country_revenue = country_sales.index[0], country_sales.iloc[0]
    top_month, top_month_revenue = monthly_sales.index[0], monthly_sales.iloc[0]
    top_product, top_product_revenue = product_sales.index[0], product_sales.iloc[0]
    country_share = top_country_revenue / net_revenue * 100 if net_revenue else 0

    report = f"""# Retail sales findings

Generated from the cleaned UCI Online Retail dataset. Values describe this dataset only.

## Key results

- **Net revenue:** £{net_revenue:,.2f} across **{len(completed_invoice_ids):,} completed orders**.
- **Average order value:** £{average_order_value:,.2f} (revenue from
  non-cancellation invoices, including their return adjustments, divided by
  completed orders).
- **Largest country by net revenue:** {top_country} at £{top_country_revenue:,.2f}
  ({country_share:.1f}% of net revenue).
- **Highest-revenue month:** {top_month} at £{top_month_revenue:,.2f}.
- **Top product by net revenue:** {top_product[1]} (stock code {top_product[0]})
  at £{top_product_revenue:,.2f}.
- **Return rows retained:** {return_rows:,}, totaling £{return_revenue:,.2f}
  in revenue. The source does not provide return reasons.
- **Cancellation rows retained:** {cancellation_rows:,}. They are excluded from
  completed-order counts and remain in net-revenue calculations.

## Suggested actions

1. Review performance by country before allocating growth effort. The largest
   country accounts for {country_share:.1f}% of net revenue; validate that
   concentration against current sales before making investment decisions.
2. Review the leading products' availability and customer feedback. The dataset
   contains no inventory or cost data, so it cannot establish stock shortages
   or product profitability.
3. Investigate negative-quantity and cancellation activity using operational
   records with return reasons; those reasons are not present in this dataset.
4. Treat monthly differences as descriptive. The dataset covers about one year
   and includes a partial December, so it is not enough to establish a reliable
   seasonal pattern or forecast.

## Data preparation summary

| Check | Rows |
|---|---:|
| Source rows | {quality["source_rows"]:,} |
| Exact duplicates removed | {quality["duplicate_rows_removed"]:,} |
| Invalid rows removed | {quality["invalid_rows_removed"]:,} |
| Clean rows retained | {quality["clean_rows"]:,} |
| Missing customer IDs retained | {quality["missing_customer_id_rows_retained"]:,} |
| Outlier rows flagged, not removed | {quality["outlier_rows_flagged_not_removed"]:,} |

## Limitations

This is historical transaction data for one retailer, not a representative
retail market sample. It contains no cost, profit, product category, inventory,
marketing, return-reason, or customer-demographic fields. Revenue must not be
presented as profit, and correlation or month-to-month changes do not establish
causation.
"""
    (REPORT_DIR / "findings.md").write_text(report, encoding="utf-8")


def main() -> None:
    if not SOURCE_FILE.exists():
        raise FileNotFoundError(
            f"Source workbook not found at {SOURCE_FILE}. "
            "Run `python src/download_data.py` first."
        )
    raw = pd.read_excel(SOURCE_FILE, engine="openpyxl")
    cleaned, quality = clean_data(raw)
    del raw
    if cleaned.empty:
        raise ValueError("Cleaning removed every transaction; check the source workbook.")
    write_outputs(cleaned, quality)
    write_findings_report(cleaned, quality)
    print(f"Prepared {quality['clean_rows']:,} rows.")
    print(f"CSV: {OUTPUT_CSV}")
    print(f"Excel: {OUTPUT_XLSX}")
    print(f"SQLite: {OUTPUT_DB}")
    print(f"Findings: {REPORT_DIR / 'findings.md'}")


if __name__ == "__main__":
    main()
